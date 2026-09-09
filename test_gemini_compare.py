"""Gemini side-by-side comparison: sanitising, match labelling, row wiring and
workbook shape. No network -- the resolver's single call site is stubbed.

The sanitiser cases are taken from the real NYFW SS27 sheet run, including the
prose answer the model gave for Wiederhoeft's IMDb column.
"""
import csv
import io
import unittest
from types import SimpleNamespace
from unittest import mock

import openpyxl

import gemini_resolver as gr
import app


class Sanitise(unittest.TestCase):
    def test_refusal_prose_becomes_blank(self):
        prose = ("I do not have enough information to answer the query. I am "
                 "missing an official verified IMDb profile or ID for the "
                 "fashion brand Wiederhoeft.")
        self.assertEqual(gr.sanitize('imdb_id', prose), '')

    def test_other_non_answers_blank(self):
        for v in ('', '   ', 'N/A', 'none', 'null', 'Unknown',
                  "I'm sorry, I could not find an official account.",
                  'No official TikTok account could be confirmed.',
                  'not applicable'):
            self.assertEqual(gr.sanitize('instagram_user', v), '', repr(v))

    def test_instagram_handle_forms(self):
        for v in ('agmesnyc', '@agmesnyc', 'https://www.instagram.com/agmesnyc',
                  'https://instagram.com/agmesnyc/', 'AGMESNYC'):
            self.assertEqual(gr.sanitize('instagram_user', v), 'agmesnyc', repr(v))

    def test_instagram_rejects_sentence(self):
        self.assertEqual(gr.sanitize('instagram_user', 'the official account is agmesnyc'), '')

    def test_twitter_handle_forms_and_limits(self):
        self.assertEqual(gr.sanitize('twitter_handle', 'https://x.com/magdabutrym'),
                         'magdabutrym')
        self.assertEqual(gr.sanitize('twitter_handle', '@zaldynyc'), 'zaldynyc')
        # >15 chars is not a valid X username
        self.assertEqual(gr.sanitize('twitter_handle', 'a' * 16), '')

    def test_facebook_must_be_a_facebook_url(self):
        self.assertEqual(gr.sanitize('facebook_page', 'http://www.facebook.com/agmesnyc'),
                         'https://www.facebook.com/agmesnyc')
        self.assertEqual(gr.sanitize('facebook_page', 'agmesnyc'), '')
        self.assertEqual(gr.sanitize('facebook_page',
                                     'https://www.facebook.com/profile.php'), '')

    def test_youtube_normalises_to_canonical_url(self):
        self.assertEqual(gr.sanitize('youtube_channel_username', '@magda_butrym'),
                         'https://www.youtube.com/@magda_butrym')
        self.assertEqual(
            gr.sanitize('youtube_channel_username',
                        'http://www.youtube.com/channel/UCwk3O80_zasUO7AD427gCcA'),
            'https://www.youtube.com/channel/UCwk3O80_zasUO7AD427gCcA')

    def test_wikipedia_keeps_language_subdomain(self):
        self.assertEqual(
            gr.sanitize('wikipedia_page', 'https://pl.wikipedia.org/wiki/Magda_Butrym'),
            'https://pl.wikipedia.org/wiki/Magda_Butrym')
        self.assertEqual(gr.sanitize('wikipedia_page', 'Magda_Butrym'), '')

    def test_imdb_extracts_id(self):
        self.assertEqual(gr.sanitize('imdb_id', 'nm2814897'), 'nm2814897')
        self.assertEqual(gr.sanitize('imdb_id', 'https://www.imdb.com/name/nm2814897/'),
                         'nm2814897')
        self.assertEqual(gr.sanitize('imdb_id', 'Zaldy'), '')


class MatchLabels(unittest.TestCase):
    def test_protocol_and_www_differences_still_match(self):
        self.assertEqual(
            app._gemini_match_label('facebook_page',
                                    'http://www.facebook.com/agmesnyc',
                                    'https://facebook.com/agmesnyc/'),
            'match')

    def test_case_and_at_sign_differences_still_match(self):
        self.assertEqual(
            app._gemini_match_label('twitter_handle', 'MagdaButrym', '@magdabutrym'),
            'match')

    def test_multi_value_cell_matches_on_any_line(self):
        existing = ('https://www.youtube.com/@JamieHaller-wt4jd\n'
                    'https://www.youtube.com/@shopjamiehaller')
        self.assertEqual(
            app._gemini_match_label('youtube_channel_username', existing,
                                    'https://www.youtube.com/@shopjamiehaller'),
            'match')

    def test_one_sided_and_blank_states(self):
        self.assertEqual(app._gemini_match_label('instagram_user', 'a', ''),
                         'existing only')
        self.assertEqual(app._gemini_match_label('instagram_user', '', 'a'),
                         'gemini only')
        self.assertEqual(app._gemini_match_label('instagram_user', '', ''),
                         'both blank')
        self.assertEqual(app._gemini_match_label('instagram_user', 'a', 'b'),
                         'mismatch')


class _FakeResp:
    """Minimal stand-in for the SDK response object _ask reads."""

    def __init__(self, text):
        self.text = text
        self.candidates = []
        self.usage_metadata = None


def _fake_generate(answers):
    """A _generate stand-in (the only network-touching function) that answers
    per entity name found in the prompt. Request accounting stays real."""
    def _gen(prompt):
        for name, payload in answers.items():
            if f"Name: {name} " in prompt:
                return _FakeResp(payload)
        return _FakeResp('')
    return _gen


class RowWiring(unittest.TestCase):
    def setUp(self):
        gr.clear_cache()
        gr.reset_stats()
        self.rows = [
            {'title': 'AGMES', 'title_category': 'Fashion',
             'title_sub_category': 'Unknown', 'instagram_user': 'agmesnyc',
             'facebook_page': '', 'twitter_handle': '', 'tiktok_user': '',
             'youtube_channel_username': '', 'wikipedia_page': '', 'imdb_id': ''},
            {'title': 'AGMES - DAR', 'title_category': 'Fashion',
             'title_sub_category': 'Unknown', 'instagram_user': 'agmesnyc',
             'facebook_page': '', 'twitter_handle': '', 'tiktok_user': '',
             'youtube_channel_username': '', 'wikipedia_page': '', 'imdb_id': ''},
        ]
        self.answer = ('{"facebook": "https://www.facebook.com/agmesnyc", '
                       '"twitter": "", "instagram": "@AGMESNYC", "youtube": "", '
                       '"tiktok": "agmesnyc", "wikipedia": "", '
                       '"imdb": "I do not have enough information."}')

    def _attach(self):
        with mock.patch.object(gr, 'available', return_value=True), \
             mock.patch.object(gr, '_generate', _fake_generate({'AGMES': self.answer})):
            return app.attach_gemini_comparison([dict(r) for r in self.rows])

    def test_disabled_resolver_leaves_rows_untouched(self):
        with mock.patch.object(gr, 'available', return_value=False):
            out = app.attach_gemini_comparison([dict(r) for r in self.rows])
        self.assertFalse(any(k.startswith('_gemini') for r in out for k in r))

    def test_values_are_sanitised_onto_underscore_keys(self):
        out = self._attach()
        self.assertEqual(out[0]['_gemini_instagram_user'], 'agmesnyc')
        self.assertEqual(out[0]['_gemini_facebook_page'],
                         'https://www.facebook.com/agmesnyc')
        # prose answer never reaches a data column
        self.assertEqual(out[0]['_gemini_imdb_id'], '')

    def test_existing_columns_are_never_modified(self):
        out = self._attach()
        for before, after in zip(self.rows, out):
            for col in app.GEMINI_COMPARE_FIELDS:
                self.assertEqual(before.get(col, ''), after.get(col, ''), col)

    def test_dar_twin_gets_identical_answers_from_cache(self):
        """The sheet gave the same entity different answers on its DAR and
        non-DAR rows; one resolve per normalised name removes that."""
        out = self._attach()
        for col in app.GEMINI_COMPARE_FIELDS:
            self.assertEqual(out[0][f'_gemini_{col}'], out[1][f'_gemini_{col}'], col)
        # both rows normalise to the same entity, so it is resolved once
        self.assertEqual(gr.stats()['requests'], 1)

    def test_repeat_resolve_is_served_from_cache(self):
        with mock.patch.object(gr, 'available', return_value=True), \
             mock.patch.object(gr, '_generate',
                               _fake_generate({'AGMES': self.answer})):
            first = gr.resolve('AGMES', 'Fashion')
            second = gr.resolve('AGMES - DAR', 'Fashion')
        self.assertEqual(first, second)
        self.assertEqual(gr.stats()['requests'], 1)
        self.assertEqual(gr.stats()['cached'], 1)

    def test_compare_records_and_summary(self):
        recs = app._gemini_compare_records(self._attach())
        self.assertEqual(len(recs), 2)
        self.assertEqual(recs[0]['instagram_user_match'], 'match')
        self.assertEqual(recs[0]['facebook_page_match'], 'gemini only')
        self.assertEqual(recs[0]['imdb_id_match'], 'both blank')
        summary = {s['field']: s for s in app._gemini_summary_records(recs)}
        self.assertEqual(summary['instagram_user']['match'], 2)
        self.assertEqual(summary['instagram_user']['agreement_when_both_filled'],
                         '100.0%')
        self.assertEqual(summary['facebook_page']['gemini_only'], 2)
        self.assertEqual(summary['imdb_id']['agreement_when_both_filled'], '')

    def test_columns_are_paired_triples(self):
        cols = app._gemini_compare_columns()
        i = cols.index('instagram_user')
        self.assertEqual(cols[i:i + 3],
                         ['instagram_user', 'instagram_user_gemini',
                          'instagram_user_match'])


class WorkbookShape(unittest.TestCase):
    def setUp(self):
        gr.clear_cache()
        gr.reset_stats()

    def _rows(self):
        base = {'title': 'AGMES', 'title_category': 'Fashion',
                'title_sub_category': 'Unknown', 'instagram_user': 'agmesnyc'}
        answer = ('{"facebook": "", "twitter": "", "instagram": "agmesnyc", '
                  '"youtube": "", "tiktok": "", "wikipedia": "", "imdb": ""}')
        with mock.patch.object(gr, 'available', return_value=True), \
             mock.patch.object(gr, '_generate', _fake_generate({'AGMES': answer})):
            return app.attach_gemini_comparison([dict(base)])

    def test_comparison_sheets_added(self):
        wb = openpyxl.load_workbook(app._rows_to_workbook(self._rows()))
        self.assertIn('Gemini Compare', wb.sheetnames)
        self.assertIn('Gemini Summary', wb.sheetnames)

    def test_ingest_sheet_keeps_its_exact_columns(self):
        wb = openpyxl.load_workbook(app._rows_to_workbook(self._rows()))
        ws = wb['Sheet1']
        header = [c.value for c in ws[1]]
        self.assertEqual(header, app.COLUMNS)
        self.assertFalse([h for h in header if h and 'gemini' in str(h)])

    def test_no_comparison_sheets_without_the_flag(self):
        plain = [{'title': 'AGMES', 'title_category': 'Fashion',
                  'instagram_user': 'agmesnyc'}]
        wb = openpyxl.load_workbook(app._rows_to_workbook(plain))
        self.assertNotIn('Gemini Compare', wb.sheetnames)


class Prompt(unittest.TestCase):
    def test_shared_rule_block_matches_the_sheet_wording(self):
        """The consolidated and per-platform prompts must carry the sheet's
        rule block unchanged -- that is what makes the comparison fair."""
        for tmpl in (gr._PER_PLATFORM_PROMPT, gr._CONSOLIDATED_PROMPT):
            self.assertIn('Never return a fan, parody, tribute, news, '
                          'aggregator, unofficial or unrelated account.', tmpl)
            self.assertIn('If no official account/page can be confidently '
                          'confirmed, return nothing.', tmpl)
            self.assertIn("please ignore the ' - DAR' suffix if present", tmpl)

    def test_consolidated_prompt_asks_for_all_seven_platforms(self):
        for label in ('Facebook', 'Twitter/X', 'Instagram', 'YouTube',
                      'TikTok', 'Wikipedia', 'IMDb'):
            self.assertIn(label, gr._CONSOLIDATED_PROMPT)

    def test_per_platform_mode_makes_one_request_per_platform(self):
        gr.clear_cache()
        gr.reset_stats()
        with mock.patch.object(gr, 'available', return_value=True), \
             mock.patch.object(gr, 'MODE', 'per_platform'), \
             mock.patch.object(gr, '_generate', lambda p: _FakeResp('')):
            gr.resolve('AGMES', 'Fashion')
        self.assertEqual(gr.stats()['requests'], len(gr.FIELDS))

    def test_consolidated_mode_makes_one_request(self):
        gr.clear_cache()
        gr.reset_stats()
        with mock.patch.object(gr, 'available', return_value=True), \
             mock.patch.object(gr, '_generate', lambda p: _FakeResp('{}')):
            gr.resolve('AGMES', 'Fashion')
        self.assertEqual(gr.stats()['requests'], 1)

    def test_request_budget_caps_a_run(self):
        gr.clear_cache()
        gr.reset_stats()
        with mock.patch.object(gr, 'available', return_value=True), \
             mock.patch.object(gr, 'MAX_REQUESTS', 2), \
             mock.patch.object(gr, '_generate', return_value=None):
            for n in range(5):
                gr.resolve(f'Brand {n}', 'Fashion')
        # 2 attempts allowed through, the rest short-circuited
        self.assertEqual(gr.stats()['requests'], 2)
        self.assertEqual(gr.stats()['errors'], 2)
        self.assertEqual(gr.stats()['capped'], 3)


class StatusEndpoint(unittest.TestCase):
    def test_status_reports_a_reason_when_unavailable(self):
        client = app.app.test_client()
        with mock.patch.object(gr, 'API_KEY', ''):
            body = client.get('/api/gemini_status').get_json()
        self.assertFalse(body['available'])
        self.assertIn('reason', body)


if __name__ == '__main__':
    unittest.main()


class EveryTitleType(unittest.TestCase):
    """The comparison must appear for every title type the UI offers, and a
    field a schema does not carry must not be scored as a Gemini-only find."""

    ANSWER = ('{"facebook":"https://www.facebook.com/x","twitter":"xhandle",'
              '"instagram":"xinsta","youtube":"@xyt","tiktok":"xtt",'
              '"wikipedia":"https://en.wikipedia.org/wiki/X","imdb":"nm1234567"}')

    KINDS = {'movie': ('Inception', 'Movies'), 'tv': ('Severance', 'TV Shows'),
             'talent': ('Tom Hanks', 'Talent'), 'game': ('Elden Ring', 'Video Games'),
             'publisher': ('Vogue', 'Publishers'), 'beauty': ('Rare Beauty', 'Beauty'),
             'beverages': ('Celsius', 'Beverages'), 'sports': ('LA Lakers', 'Sports Teams'),
             'general': ('Peloton', 'General')}

    def _generate(self, titles_type):
        titles = list(titles_type)
        payload = {'titles': titles, 'includeDar': False, 'autoFetch': False,
                   'geminiCompare': True, 'titles_type': titles_type}
        gr.clear_cache()
        gr.reset_stats()
        with mock.patch.object(gr, 'available', return_value=True), \
             mock.patch.object(gr, '_generate',
                               lambda p: _FakeResp(self.ANSWER)):
            resp = app.app.test_client().post('/api/generate', json=payload)
        self.assertEqual(resp.status_code, 200)
        return openpyxl.load_workbook(io.BytesIO(resp.data))

    def test_each_title_type_gets_the_comparison_sheets(self):
        for kind, (title, label) in self.KINDS.items():
            wb = self._generate({title: kind})
            self.assertIn('Gemini Compare', wb.sheetnames, kind)
            self.assertIn('Gemini Summary', wb.sheetnames, kind)
            ws = wb['Gemini Compare']
            hdr = [c.value for c in ws[1]]
            self.assertEqual(ws.max_row - 1, 1, kind)
            self.assertEqual(ws.cell(row=2, column=hdr.index('title_type') + 1).value,
                             label, kind)

    def test_mixed_run_labels_each_row_with_its_type(self):
        wb = self._generate({'Inception': 'movie', 'Tom Hanks': 'talent',
                             'Elden Ring': 'game'})
        ws = wb['Gemini Compare']
        hdr = [c.value for c in ws[1]]
        col = hdr.index('title_type') + 1
        types = {ws.cell(row=r, column=col).value for r in range(2, ws.max_row + 1)}
        self.assertEqual(types, {'Movies', 'Talent', 'Video Games'})

    def test_mixed_run_summary_is_per_type_plus_a_combined_block(self):
        wb = self._generate({'Inception': 'movie', 'Tom Hanks': 'talent'})
        ws = wb['Gemini Summary']
        hdr = [c.value for c in ws[1]]
        col = hdr.index('title_type') + 1
        labels = [ws.cell(row=r, column=col).value for r in range(2, ws.max_row + 1)]
        self.assertEqual(set(labels), {'Movies', 'Talent', '(all types)'})
        # one row per field per block
        self.assertEqual(len(labels), 3 * len(app.GEMINI_COMPARE_FIELDS))

    def test_field_absent_from_a_schema_is_not_scored(self):
        """Beauty and Beverages carry no imdb_id column."""
        for kind, title in (('beauty', 'Rare Beauty'), ('beverages', 'Celsius')):
            wb = self._generate({title: kind})
            ws = wb['Gemini Compare']
            hdr = [c.value for c in ws[1]]
            self.assertEqual(
                ws.cell(row=2, column=hdr.index('imdb_id_match') + 1).value,
                'not in schema', kind)
            summary = {(r[hdr2.index('title_type')], r[hdr2.index('field')]): r
                       for hdr2, r in [([c.value for c in wb['Gemini Summary'][1]],
                                        [c.value for c in row])
                                       for row in wb['Gemini Summary'].iter_rows(min_row=2)]}
            key = [k for k in summary if k[1] == 'imdb_id'][0]
            row = summary[key]
            h = [c.value for c in wb['Gemini Summary'][1]]
            self.assertEqual(row[h.index('gemini_only')], 0, kind)
            self.assertEqual(row[h.index('not_in_schema')], 1, kind)

    def test_imdb_is_still_scored_where_the_schema_has_it(self):
        wb = self._generate({'Inception': 'movie'})
        ws = wb['Gemini Compare']
        hdr = [c.value for c in ws[1]]
        self.assertEqual(ws.cell(row=2, column=hdr.index('imdb_id_match') + 1).value,
                         'gemini only')


class _FakeHTTP:
    """Stand-in for requests.post: records payloads, replays canned bodies."""

    def __init__(self, body, status=200):
        self.body = body
        self.status = status
        self.calls = []

    def __call__(self, url, json=None, timeout=None, allow_redirects=None):
        self.calls.append({'url': url, 'json': json})
        outer = self

        class R:
            status_code = outer.status

            def raise_for_status(self):
                if outer.status >= 400:
                    raise RuntimeError('HTTP %s' % outer.status)

            def json(self):
                return outer.body() if callable(outer.body) else outer.body

        return R()


class AppsScriptSource(unittest.TestCase):
    """The Apps Script bridge: native Sheets =GEMINI(), no API quota."""

    URL = 'https://script.google.com/macros/s/EXAMPLE/exec'
    TOKEN = 'shared-secret'

    def setUp(self):
        gr.clear_cache()
        gr.reset_stats()

    def _as(self, **extra):
        patches = {'SOURCE': 'appsscript', 'SCRIPT_URL': self.URL,
                   'SCRIPT_TOKEN': self.TOKEN}
        patches.update(extra)
        return [mock.patch.object(gr, k, v) for k, v in patches.items()]

    def _run(self, fake, fn, **extra):
        stack = self._as(**extra)
        for p in stack:
            p.start()
        try:
            with mock.patch.dict('sys.modules'):
                with mock.patch('requests.post', fake):
                    return fn()
        finally:
            for p in reversed(stack):
                p.stop()

    def test_availability_needs_url_and_token(self):
        with mock.patch.object(gr, 'SOURCE', 'appsscript'), \
             mock.patch.object(gr, 'SCRIPT_URL', ''), \
             mock.patch.object(gr, 'SCRIPT_TOKEN', self.TOKEN):
            self.assertFalse(gr.available())
        with mock.patch.object(gr, 'SOURCE', 'appsscript'), \
             mock.patch.object(gr, 'SCRIPT_URL', self.URL), \
             mock.patch.object(gr, 'SCRIPT_TOKEN', ''):
            self.assertFalse(gr.available())
        with mock.patch.object(gr, 'SOURCE', 'appsscript'), \
             mock.patch.object(gr, 'SCRIPT_URL', self.URL), \
             mock.patch.object(gr, 'SCRIPT_TOKEN', self.TOKEN):
            self.assertTrue(gr.available())

    def test_availability_does_not_need_an_api_key(self):
        """The whole point: no GEMINI_API_KEY, still available."""
        with mock.patch.object(gr, 'SOURCE', 'appsscript'), \
             mock.patch.object(gr, 'API_KEY', ''), \
             mock.patch.object(gr, 'SCRIPT_URL', self.URL), \
             mock.patch.object(gr, 'SCRIPT_TOKEN', self.TOKEN):
            self.assertTrue(gr.available())
            st = gr.status()
        self.assertEqual(st['source'], 'appsscript')
        self.assertFalse(st['metered'])

    def test_api_source_reports_metered(self):
        with mock.patch.object(gr, 'SOURCE', 'api'):
            self.assertTrue(gr.status()['metered'])

    def test_values_are_sanitised_like_the_api_path(self):
        fake = _FakeHTTP({'results': [{
            'name': 'Tom Hanks', 'context': 'Talent',
            'facebook': 'http://www.facebook.com/TomHanks',
            'twitter': '@tomhanks', 'instagram': 'TOMHANKS',
            'youtube': '', 'tiktok': '',
            'wikipedia': 'https://en.wikipedia.org/wiki/Tom_Hanks',
            'imdb': 'I do not have enough information.'}]})
        out = self._run(fake, lambda: gr.resolve('Tom Hanks - DAR', 'Talent'))
        self.assertEqual(out['facebook_page'], 'https://www.facebook.com/TomHanks')
        self.assertEqual(out['twitter_handle'], 'tomhanks')
        self.assertEqual(out['instagram_user'], 'tomhanks')
        self.assertEqual(out['wikipedia_page'], 'https://en.wikipedia.org/wiki/Tom_Hanks')
        self.assertEqual(out['imdb_id'], '')          # prose rejected here too

    def test_dar_suffix_and_token_are_sent_correctly(self):
        fake = _FakeHTTP({'results': []})
        self._run(fake, lambda: gr.resolve('Tom Hanks - DAR', 'Talent'))
        self.assertEqual(len(fake.calls), 1)
        sent = fake.calls[0]['json']
        self.assertEqual(sent['token'], self.TOKEN)
        self.assertEqual(sent['entities'], [{'name': 'Tom Hanks', 'context': 'Talent'}])
        self.assertEqual(fake.calls[0]['url'], self.URL)

    def test_reply_is_matched_by_name_not_order(self):
        fake = _FakeHTTP({'results': [
            {'name': 'Zendaya', 'instagram': 'zendaya'},
            {'name': 'Tom Hanks', 'instagram': 'tomhanks'},
        ]})
        got = self._run(fake, lambda: gr.resolve_many(
            [('Tom Hanks', 'Talent'), ('Zendaya', 'Talent')]))
        self.assertEqual(got[('Tom Hanks', 'Talent')]['instagram_user'], 'tomhanks')
        self.assertEqual(got[('Zendaya', 'Talent')]['instagram_user'], 'zendaya')

    def test_entities_are_batched_not_one_request_each(self):
        fake = _FakeHTTP({'results': []})
        items = [('Brand %d' % i, 'Fashion') for i in range(7)]
        self._run(fake, lambda: gr.resolve_many(items), SCRIPT_BATCH=3, WORKERS=1)
        self.assertEqual(len(fake.calls), 3)          # 3 + 3 + 1
        sizes = sorted(len(c['json']['entities']) for c in fake.calls)
        self.assertEqual(sizes, [1, 3, 3])

    def test_dar_twins_collapse_into_one_entity(self):
        fake = _FakeHTTP({'results': [{'name': 'AGMES', 'instagram': 'agmesnyc'}]})
        got = self._run(fake, lambda: gr.resolve_many(
            [('AGMES', 'Fashion'), ('AGMES - DAR', 'Fashion')]), WORKERS=1)
        self.assertEqual(len(fake.calls), 1)
        self.assertEqual(len(fake.calls[0]['json']['entities']), 1)
        self.assertEqual(got[('AGMES', 'Fashion')]['instagram_user'], 'agmesnyc')
        self.assertEqual(got[('AGMES - DAR', 'Fashion')]['instagram_user'], 'agmesnyc')

    def test_progress_counts_every_original_pair(self):
        fake = _FakeHTTP({'results': []})
        seen = []
        self._run(fake, lambda: gr.resolve_many(
            [('A', 'X'), ('A - DAR', 'X'), ('B', 'X')],
            progress=lambda d, t: seen.append((d, t))), WORKERS=1)
        self.assertTrue(seen)
        self.assertEqual(seen[-1], (3, 3))

    def test_error_body_yields_blanks_not_an_exception(self):
        fake = _FakeHTTP({'error': 'unauthorized'})
        out = self._run(fake, lambda: gr.resolve('Tom Hanks', 'Talent'))
        self.assertEqual(out, {f: '' for f in gr.FIELDS})
        self.assertEqual(gr.stats()['errors'], 1)

    def test_http_failure_yields_blanks_not_an_exception(self):
        fake = _FakeHTTP({}, status=500)
        out = self._run(fake, lambda: gr.resolve('Tom Hanks', 'Talent'))
        self.assertEqual(out, {f: '' for f in gr.FIELDS})
        self.assertEqual(gr.stats()['errors'], 1)

    def test_no_search_cost_is_attributed_to_this_source(self):
        fake = _FakeHTTP({'results': [{'name': 'Tom Hanks', 'instagram': 'tomhanks'}]})
        self._run(fake, lambda: gr.resolve('Tom Hanks', 'Talent'))
        s = gr.stats()
        self.assertEqual(s['grounded'], 0)
        self.assertEqual(s['est_search_cost_usd'], 0)

    def test_end_to_end_export_uses_the_script_source(self):
        fake = _FakeHTTP({'results': [{'name': 'Tom Hanks', 'instagram': 'tomhanks',
                                       'imdb': 'nm0000158'}]})

        def _go():
            return app.app.test_client().post('/api/generate', json={
                'titles': ['Tom Hanks'], 'includeDar': False, 'autoFetch': False,
                'geminiCompare': True, 'titles_type': {'Tom Hanks': 'talent'}})

        resp = self._run(fake, _go)
        self.assertEqual(resp.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(resp.data))
        self.assertIn('Gemini Compare', wb.sheetnames)
        ws = wb['Gemini Compare']
        hdr = [c.value for c in ws[1]]
        self.assertEqual(
            ws.cell(row=2, column=hdr.index('instagram_user_gemini') + 1).value,
            'tomhanks')
        self.assertEqual(
            ws.cell(row=2, column=hdr.index('imdb_id_gemini') + 1).value,
            'nm0000158')

    def test_status_endpoint_describes_the_script_source(self):
        stack = self._as()
        for p in stack:
            p.start()
        try:
            body = app.app.test_client().get('/api/gemini_status').get_json()
        finally:
            for p in reversed(stack):
                p.stop()
        self.assertTrue(body['available'])
        self.assertEqual(body['source'], 'appsscript')
        self.assertFalse(body['metered'])


class SheetSource(unittest.TestCase):
    """GEMINI_SOURCE=sheet: the three-phase, Workspace-covered flow."""

    URL = 'https://script.google.com/macros/s/EXAMPLE/exec'
    TOKEN = 'shared-secret'

    def setUp(self):
        gr.clear_cache()
        gr.reset_stats()

    def _on(self, **extra):
        base = {'SOURCE': 'sheet', 'SCRIPT_URL': self.URL, 'SCRIPT_TOKEN': self.TOKEN}
        base.update(extra)
        return [mock.patch.object(gr, k, v) for k, v in base.items()]

    def _run(self, fake, fn, **extra):
        stack = self._on(**extra)
        for p in stack:
            p.start()
        try:
            with mock.patch('requests.post', fake):
                return fn()
        finally:
            for p in reversed(stack):
                p.stop()

    def test_source_is_interactive_and_available_without_a_key(self):
        stack = self._on(API_KEY='')
        for p in stack:
            p.start()
        try:
            self.assertTrue(gr.available())
            self.assertTrue(gr.interactive())
            st = gr.status()
        finally:
            for p in reversed(stack):
                p.stop()
        self.assertEqual(st['source'], 'sheet')
        self.assertTrue(st['interactive'])
        self.assertFalse(st['metered'])

    def test_synchronous_resolve_returns_nothing_for_this_source(self):
        """There is no synchronous answer -- the flow needs a human step, and
        pretending otherwise would silently produce blank columns."""
        stack = self._on()
        for p in stack:
            p.start()
        try:
            self.assertEqual(gr.resolve('Tom Hanks', 'Talent'), {})
            self.assertEqual(gr.resolve_many([('Tom Hanks', 'Talent')]), {})
        finally:
            for p in reversed(stack):
                p.stop()

    def test_push_sends_action_token_and_existing_values(self):
        fake = _FakeHTTP({'batch': 'cmp_1', 'url': 'https://x/#gid=1', 'rows': 1})
        res = self._run(fake, lambda: gr.sheet_push([{
            'title': 'Tom Hanks', 'title_type': 'Talent', 'context': 'Talent',
            'existing': {'instagram_user': 'tomhanks'}}]))
        sent = fake.calls[0]['json']
        self.assertEqual(sent['action'], 'push')
        self.assertEqual(sent['token'], self.TOKEN)
        self.assertEqual(sent['rows'][0]['existing']['instagram_user'], 'tomhanks')
        self.assertEqual(sent['rows'][0]['existing']['imdb_id'], '')
        self.assertEqual(res['batch'], 'cmp_1')

    def test_activate_and_status_pass_the_batch(self):
        fake = _FakeHTTP({'batch': 'cmp_1', 'activated': 14})
        self._run(fake, lambda: gr.sheet_activate('cmp_1'))
        self.assertEqual(fake.calls[0]['json']['action'], 'activate')
        self.assertEqual(fake.calls[0]['json']['batch'], 'cmp_1')

        fake2 = _FakeHTTP({'status': {'filled': 9, 'text': 0, 'pending': 5}})
        st = self._run(fake2, lambda: gr.sheet_status('cmp_1'))
        self.assertEqual(fake2.calls[0]['json']['action'], 'status')
        self.assertEqual(st['filled'], 9)

    def test_pull_sanitises_and_keeps_existing_values(self):
        fake = _FakeHTTP({'status': {'filled': 2},
                          'results': [{
                              'title': 'Tom Hanks', 'title_type': 'Talent',
                              'context': 'Talent / Actor',
                              'existing': {'instagram_user': 'tomhanks'},
                              'instagram_user': '@TOMHANKS',
                              'imdb_id': 'I do not have enough information.'}]})
        status, rows = self._run(fake, lambda: gr.sheet_pull('cmp_1'))
        self.assertEqual(status['filled'], 2)
        self.assertEqual(rows[0]['instagram_user'], 'tomhanks')
        self.assertEqual(rows[0]['imdb_id'], '')
        self.assertEqual(rows[0]['existing']['instagram_user'], 'tomhanks')

    def test_script_errors_surface_rather_than_being_swallowed(self):
        fake = _FakeHTTP({'error': 'unauthorized'})
        stack = self._on()
        for p in stack:
            p.start()
        try:
            with mock.patch('requests.post', fake):
                with self.assertRaises(gr.ScriptError):
                    gr.sheet_activate('cmp_1')
        finally:
            for p in reversed(stack):
                p.stop()

    def test_scoring_from_a_pulled_batch_is_schema_aware(self):
        pulled = [
            {'title': 'Tom Hanks', 'title_type': 'Talent', 'context': 'Talent',
             'existing': {'instagram_user': 'tomhanks'},
             'instagram_user': 'tomhanks', 'imdb_id': 'nm0000158'},
            {'title': 'Rare Beauty', 'title_type': 'Beauty', 'context': 'Beauty',
             'existing': {'instagram_user': ''},
             'instagram_user': 'rarebeauty', 'imdb_id': 'nm9999999'},
        ]
        recs = app.gemini_records_from_pull(pulled)
        self.assertEqual(recs[0]['instagram_user_match'], 'match')
        self.assertEqual(recs[0]['imdb_id_match'], 'gemini only')
        self.assertEqual(recs[1]['instagram_user_match'], 'gemini only')
        # Beauty has no imdb_id column, so it must not be scored as a find
        self.assertEqual(recs[1]['imdb_id_match'], 'not in schema')
        summary = {(r['title_type'], r['field']): r
                   for r in app._gemini_summary_records(recs)}
        self.assertEqual(summary[('Beauty', 'imdb_id')]['not_in_schema'], 1)
        self.assertEqual(summary[('Beauty', 'imdb_id')]['gemini_only'], 0)

    def test_pull_endpoint_returns_scored_rows_and_summary(self):
        fake = _FakeHTTP({'status': {'filled': 1, 'text': 0, 'pending': 6},
                          'results': [{
                              'title': 'Tom Hanks', 'title_type': 'Talent',
                              'context': 'Talent',
                              'existing': {'instagram_user': 'tomhanks'},
                              'instagram_user': 'tomhanks'}]})
        body = self._run(fake, lambda: app.app.test_client().post(
            '/api/gemini_sheet/pull', json={'batch': 'cmp_1'}).get_json())
        self.assertEqual(body['batch'], 'cmp_1')
        self.assertEqual(body['rows'][0]['instagram_user_match'], 'match')
        self.assertIn('title_type', body['columns'])
        self.assertTrue(body['summary'])

    def test_download_endpoint_builds_the_comparison_sheets(self):
        fake = _FakeHTTP({'status': {}, 'results': [{
            'title': 'Tom Hanks', 'title_type': 'Talent', 'context': 'Talent',
            'existing': {'instagram_user': 'tomhanks'},
            'instagram_user': 'tomhanks'}]})
        resp = self._run(fake, lambda: app.app.test_client().get(
            '/api/gemini_sheet/download?batch=cmp_1'))
        self.assertEqual(resp.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(resp.data))
        self.assertEqual(wb.sheetnames,
                         ['Gemini Compare', 'Gemini Summary', 'Gemini Run Notes'])

    def test_endpoints_require_a_batch(self):
        c = app.app.test_client()
        self.assertEqual(c.post('/api/gemini_sheet/activate', json={}).status_code, 400)
        self.assertEqual(c.get('/api/gemini_sheet/status').status_code, 400)
        self.assertEqual(c.post('/api/gemini_sheet/pull', json={}).status_code, 400)

    def test_push_endpoint_refuses_when_unconfigured(self):
        with mock.patch.object(gr, 'SOURCE', 'sheet'), \
             mock.patch.object(gr, 'SCRIPT_URL', ''), \
             mock.patch.object(gr, 'SCRIPT_TOKEN', ''):
            r = app.app.test_client().post('/api/gemini_sheet/push',
                                           json={'titles': ['Inception']})
        self.assertEqual(r.status_code, 400)


class RunNotes(unittest.TestCase):
    """A run where every call failed must not look like a run that found
    nothing. Reproduces the 429 RESOURCE_EXHAUSTED case seen in production,
    where a free-tier key returned quota errors for every request and the
    export came back full of 'existing only' with no indication why."""

    QUOTA = ("429 RESOURCE_EXHAUSTED. {'error': {'code': 429, 'message': "
             "'You exceeded your current quota, please check your plan and "
             "billing details.'}}")

    def setUp(self):
        gr.clear_cache()
        gr.reset_stats()

    def _notes(self, patch):
        with mock.patch.object(gr, 'available', return_value=True), patch:
            resp = app.app.test_client().post('/api/generate', json={
                'titles': ['Mitchell Starc'], 'includeDar': False,
                'autoFetch': False, 'geminiCompare': True,
                'titles_type': {'Mitchell Starc': 'talent'}})
        self.assertEqual(resp.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(resp.data))
        self.assertIn('Gemini Run Notes', wb.sheetnames)
        return wb, {r[0]: r[1] for r in
                    wb['Gemini Run Notes'].iter_rows(min_row=2, values_only=True)}

    def _quota_client(self):
        def boom(*a, **k):
            raise RuntimeError(self.QUOTA)
        return mock.patch.object(gr, '_get_client', boom)

    def test_total_failure_is_reported_as_failed(self):
        _wb, d = self._notes(self._quota_client())
        self.assertEqual(d['status'], 'FAILED')
        self.assertIn('NOT a "nothing found" result', d['what this means'])
        self.assertEqual(d['requests failed'], 1)
        self.assertEqual(d['gemini values returned'], 0)

    def test_the_google_error_text_is_carried_into_the_workbook(self):
        _wb, d = self._notes(self._quota_client())
        self.assertIn('RESOURCE_EXHAUSTED', d['last error from Google'])

    def test_a_genuine_empty_result_is_not_called_a_failure(self):
        patch = mock.patch.object(gr, '_generate',
                                  lambda p: _FakeResp('{"instagram":""}'))
        _wb, d = self._notes(patch)
        self.assertEqual(d['status'], 'EMPTY')
        self.assertEqual(d['requests failed'], 0)
        self.assertIn('real "nothing found"', d['what this means'])

    def test_a_good_run_is_reported_ok(self):
        patch = mock.patch.object(
            gr, '_generate',
            lambda p: _FakeResp('{"instagram":"mstarc56","imdb":"nm10052216"}'))
        _wb, d = self._notes(patch)
        self.assertEqual(d['status'], 'OK')
        self.assertEqual(d['gemini values returned'], 2)

    def test_errors_are_counted_even_when_raised_outside_the_sdk_call(self):
        """The outer per-entity handler used to swallow exceptions without
        counting them, which is how a failed run reported errors: 0."""
        with mock.patch.object(gr, 'available', return_value=True), \
             mock.patch.object(gr, '_resolve_one_api',
                               mock.Mock(side_effect=RuntimeError('boom'))):
            gr.resolve_many([('X', 'Talent')])
        self.assertEqual(gr.stats()['errors'], 1)
        self.assertIn('boom', gr.stats()['last_error'])

    def test_notes_appear_on_the_sheet_source_download_too(self):
        fake = _FakeHTTP({'status': {}, 'results': [{
            'title': 'Tom Hanks', 'title_type': 'Talent', 'context': 'Talent',
            'existing': {'instagram_user': 'tomhanks'},
            'instagram_user': 'tomhanks'}]})
        stack = [mock.patch.object(gr, k, v) for k, v in
                 (('SOURCE', 'sheet'),
                  ('SCRIPT_URL', 'https://script.google.com/macros/s/E/exec'),
                  ('SCRIPT_TOKEN', 't'))]
        for p in stack:
            p.start()
        try:
            with mock.patch('requests.post', fake):
                resp = app.app.test_client().get(
                    '/api/gemini_sheet/download?batch=cmp_1')
        finally:
            for p in reversed(stack):
                p.stop()
        wb = openpyxl.load_workbook(io.BytesIO(resp.data))
        self.assertEqual(wb.sheetnames,
                         ['Gemini Compare', 'Gemini Summary', 'Gemini Run Notes'])


class ProbeAndReporting(unittest.TestCase):
    """The probe exists to tell two indistinguishable 429s apart, and the run
    notes exist to explain an empty column. Both are diagnostics, so a wrong
    answer here is worse than no answer."""

    def setUp(self):
        self._sdk, self._key = gr.SDK_OK, gr.API_KEY
        gr.SDK_OK, gr.API_KEY = True, 'test-key'
        gr.reset_stats()

    def tearDown(self):
        gr.SDK_OK, gr.API_KEY = self._sdk, self._key
        gr.reset_stats()

    @staticmethod
    def _client(behaviour):
        """behaviour(grounded) -> reply text, or raises."""
        class Models:
            def generate_content(self, model, contents, config):
                grounded = bool(getattr(config, 'tools', None))
                text = behaviour(grounded)
                return SimpleNamespace(text=text)
        return SimpleNamespace(models=Models())

    def test_no_key_reports_plainly(self):
        gr.API_KEY = ''
        r = gr.probe()
        self.assertFalse(r['ok'])
        self.assertIn('no API key', r['reason'])

    def test_both_work(self):
        with mock.patch.object(gr, '_get_client',
                               return_value=self._client(lambda g: 'OK')):
            r = gr.probe()
        self.assertTrue(r['ok'])
        self.assertTrue(r['plain']['ok'])
        self.assertTrue(r['grounded']['ok'])
        self.assertIn('Both work', r['verdict'])

    def test_grounding_is_the_blocker(self):
        def behaviour(grounded):
            if grounded:
                raise RuntimeError("429 RESOURCE_EXHAUSTED {'quotaMetric': "
                                   "'generativelanguage.googleapis.com/grounded_search'}")
            return 'OK'
        with mock.patch.object(gr, '_get_client',
                               return_value=self._client(behaviour)):
            r = gr.probe()
        self.assertTrue(r['plain']['ok'])
        self.assertFalse(r['grounded']['ok'])
        self.assertEqual(r['grounded']['quota'],
                         'generativelanguage.googleapis.com/grounded_search')
        self.assertIn('NOT with Google Search grounding', r['verdict'])

    def test_key_cannot_call_at_all(self):
        def behaviour(grounded):
            raise RuntimeError("429 RESOURCE_EXHAUSTED {'quotaMetric': "
                               "'generativelanguage.googleapis.com/"
                               "generate_content_free_tier_requests'}")
        with mock.patch.object(gr, '_get_client',
                               return_value=self._client(behaviour)):
            r = gr.probe()
        self.assertFalse(r['ok'])
        self.assertIn('cannot call this model at all', r['verdict'])
        self.assertIn('free_tier_requests', r['plain']['quota'])

    def test_probe_does_not_spend_the_run_budget(self):
        with mock.patch.object(gr, '_get_client',
                               return_value=self._client(lambda g: 'OK')):
            gr.probe()
        st = gr.stats()
        self.assertEqual(st['requests'], 0, 'probe must not count as run requests')
        self.assertEqual(st['errors'], 0)

    def test_stats_report_the_real_config_not_defaults(self):
        st = gr.stats()
        self.assertEqual(st['source'], gr.SOURCE)
        self.assertEqual(st['model'], gr.MODEL)
        self.assertEqual(st['mode'], gr.MODE)

    def test_quota_name_is_lifted_to_the_front_of_the_error(self):
        gr._note_error(RuntimeError(
            "429 RESOURCE_EXHAUSTED. " + ("boilerplate " * 40)
            + "{'quotaMetric': 'generativelanguage.googleapis.com/"
              "generate_content_free_tier_requests'}"))
        err = gr.stats()['last_error']
        self.assertTrue(err.startswith('Quota exhausted: '), err[:60])
        self.assertIn('free_tier_requests', err[:120],
                      'the quota name must survive truncation')

    def test_non_quota_error_gets_no_quota_prefix(self):
        gr._note_error(RuntimeError('404 NOT_FOUND: unknown model'))
        self.assertEqual(gr.stats()['last_error'], '404 NOT_FOUND: unknown model')


class SheetFileRoundTrip(unittest.TestCase):
    """The no-key route: export questions as formulas, score answers back.

    This path must work on a deployment where the metered API is refused, so
    every test here runs with the resolver deliberately unavailable.
    """

    def setUp(self):
        self._key, self._sdk = gr.API_KEY, gr.SDK_OK
        gr.API_KEY, gr.SDK_OK = '', False          # no key: the whole point
        self.client = app.app.test_client()

    def tearDown(self):
        gr.API_KEY, gr.SDK_OK = self._key, self._sdk

    # -- export -------------------------------------------------------------

    def test_export_works_with_no_api_key(self):
        self.assertFalse(gr.available(), 'precondition: resolver unavailable')
        r = self.client.post('/api/gemini_sheet/export', json={
            'titles': ['Tom Hanks'], 'titles_type': {'Tom Hanks': 'talent'},
            'includeDar': False})
        self.assertEqual(r.status_code, 200)
        self.assertIn('text/csv', r.headers['Content-Type'])
        body = r.data.decode('utf-8-sig')
        self.assertIn('instagram_user_gemini', body)
        # the raw body is CSV-quoted, so the formula's own quotes are doubled
        self.assertIn('=GEMINI(""Using Google Search', body)
        parsed = list(csv.DictReader(io.StringIO(body)))
        self.assertTrue(parsed[0]['instagram_user_gemini'].startswith('=GEMINI("'))

    def test_export_asks_only_for_fields_in_the_row_schema(self):
        rows = [{'title': 'X', 'title_category': 'Talent', '_tfx_schema': 'talent'}]
        recs = app.gemini_sheet_export_rows(rows)
        schema = app._gemini_row_schema_columns(rows[0])
        for f in app.GEMINI_COMPARE_FIELDS:
            has_formula = recs[0][f + '_gemini'].startswith('=GEMINI(')
            self.assertEqual(has_formula, f in schema,
                             '%s: formula=%s in-schema=%s' % (f, has_formula, f in schema))

    def test_export_csv_is_parseable_and_keeps_formulas_intact(self):
        rows = [{'title': 'Dwayne "The Rock" Johnson', 'title_category': 'Talent',
                 '_tfx_schema': 'talent'}]
        data, n = app.gemini_sheet_export_csv(rows)
        self.assertEqual(n, 1)
        back = list(csv.DictReader(io.StringIO(data.decode('utf-8-sig'))))
        self.assertEqual(len(back), 1)
        f = back[0]['instagram_user_gemini']
        self.assertTrue(f.startswith('=GEMINI("'), f[:40])
        self.assertTrue(f.endswith('")'), f[-20:])
        self.assertIn('""The Rock""', f, 'inner quotes must stay doubled')

    # -- score --------------------------------------------------------------

    @staticmethod
    def _csv(rows, headers):
        buf = io.StringIO()
        w = csv.DictWriter(buf, fieldnames=headers, lineterminator='\n')
        w.writeheader()
        for r in rows:
            w.writerow(r)
        return (io.BytesIO(buf.getvalue().encode()), 'done.csv')

    HDR = ['title', 'title_type', 'title_category', 'context_sent_to_gemini',
           'instagram_user', 'instagram_user_gemini']

    def _score(self, rows, fmt='json'):
        r = self.client.post('/api/gemini_sheet/score', data={
            'file': self._csv(rows, self.HDR), 'format': fmt},
            content_type='multipart/form-data')
        return r

    def test_scores_a_generated_sheet(self):
        r = self._score([
            {'title': 'A', 'title_type': 'Talent', 'instagram_user': 'a',
             'instagram_user_gemini': 'a'},
            {'title': 'B', 'title_type': 'Talent', 'instagram_user': 'b',
             'instagram_user_gemini': 'other'},
            {'title': 'C', 'title_type': 'Talent', 'instagram_user': '',
             'instagram_user_gemini': 'c'},
        ])
        self.assertEqual(r.status_code, 200)
        got = [x['instagram_user_match'] for x in r.get_json()['rows']]
        self.assertEqual(got, ['match', 'mismatch', 'gemini only'])
        self.assertEqual(r.get_json()['notes'][0]['value'], 'OK')

    def test_ungenerated_formula_is_not_scored_as_nothing_found(self):
        r = self._score([
            {'title': 'A', 'title_type': 'Talent', 'instagram_user': 'a',
             'instagram_user_gemini': '=GEMINI("Using Google Search, ...")'},
        ])
        j = r.get_json()
        self.assertEqual(j['rows'][0]['instagram_user_match'], 'not generated')
        self.assertEqual(j['rows'][0]['instagram_user_gemini'], '')
        self.assertEqual(j['pending'], 1)
        self.assertEqual(j['notes'][0]['value'], 'NOT GENERATED')
        self.assertIn('Generate and fill', j['notes'][1]['value'])

    def test_error_cell_is_also_treated_as_not_generated(self):
        r = self._score([{'title': 'A', 'title_type': 'Talent',
                          'instagram_user': 'a', 'instagram_user_gemini': '#ERROR!'}])
        self.assertEqual(r.get_json()['rows'][0]['instagram_user_match'],
                         'not generated')

    def test_partial_generation_is_called_partial(self):
        r = self._score([
            {'title': 'A', 'title_type': 'Talent', 'instagram_user': 'a',
             'instagram_user_gemini': 'a'},
            {'title': 'B', 'title_type': 'Talent', 'instagram_user': 'b',
             'instagram_user_gemini': '=GEMINI("x")'},
        ])
        j = r.get_json()
        self.assertEqual(j['notes'][0]['value'], 'PARTIAL')
        self.assertEqual(j['pending'], 1)

    def test_all_generated_all_empty_is_a_real_nothing_found(self):
        r = self._score([{'title': 'A', 'title_type': 'Talent',
                          'instagram_user': 'a', 'instagram_user_gemini': ''}])
        j = r.get_json()
        self.assertEqual(j['notes'][0]['value'], 'EMPTY')
        self.assertEqual(j['rows'][0]['instagram_user_match'], 'existing only')

    def test_rejects_a_file_with_no_gemini_columns(self):
        r = self.client.post('/api/gemini_sheet/score', data={
            'file': self._csv([{'title': 'A'}], ['title'])},
            content_type='multipart/form-data')
        self.assertEqual(r.status_code, 400)
        self.assertIn('_gemini', r.get_json()['error'])

    def test_rejects_a_missing_file(self):
        r = self.client.post('/api/gemini_sheet/score', data={},
                             content_type='multipart/form-data')
        self.assertEqual(r.status_code, 400)

    def test_scored_download_is_a_three_sheet_workbook(self):
        r = self._score([{'title': 'A', 'title_type': 'Talent',
                          'instagram_user': 'a', 'instagram_user_gemini': 'a'}],
                        fmt='xlsx')
        self.assertEqual(r.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(r.data))
        self.assertEqual(wb.sheetnames,
                         ['Gemini Compare', 'Gemini Summary', 'Gemini Run Notes'])

    # -- the whole loop -----------------------------------------------------

    def test_export_then_generate_then_score(self):
        """Export, simulate a person pressing Generate and fill, score it."""
        rows = [{'title': 'Tom Hanks - DAR', 'title_category': 'Talent',
                 '_tfx_schema': 'talent', 'instagram_user': ''}]
        data, _ = app.gemini_sheet_export_csv(rows)
        parsed = list(csv.DictReader(io.StringIO(data.decode('utf-8-sig'))))
        self.assertTrue(parsed[0]['instagram_user_gemini'].startswith('=GEMINI('))

        # what Sheets does when the button is pressed: EVERY formula in the
        # selection resolves, most of them to nothing for this entity
        for f in app.GEMINI_COMPARE_FIELDS:
            if parsed[0][f + '_gemini'].startswith('=GEMINI('):
                parsed[0][f + '_gemini'] = 'tomhanks' if f == 'instagram_user' else ''

        records, pending = app.gemini_score_uploaded(parsed)
        self.assertEqual(pending, 0)
        self.assertEqual(records[0]['instagram_user_match'], 'gemini only')
        self.assertEqual(records[0]['title'], 'Tom Hanks - DAR',
                         'the DAR suffix belongs in the sheet, only the prompt drops it')


class PerRunBudget(unittest.TestCase):
    """GEMINI_MAX_REQUESTS has to mean "this run", because that is what it is
    called and what the run notes claim. Before this, the counter accumulated
    for the life of the process: past the limit every further title came back
    blank and was reported as capped rather than as an error, which on a
    metered key is a silent data-quality failure rather than a billing one."""

    def setUp(self):
        self._key, self._sdk, self._cap = gr.API_KEY, gr.SDK_OK, gr.MAX_REQUESTS
        gr.API_KEY, gr.SDK_OK = 'k', True
        gr.reset_stats()
        gr.clear_cache()

    def tearDown(self):
        gr.API_KEY, gr.SDK_OK, gr.MAX_REQUESTS = self._key, self._sdk, self._cap
        gr.reset_stats()
        gr.clear_cache()

    @staticmethod
    def _rows(names):
        return [{'title': n, 'title_category': 'Talent', '_tfx_schema': 'talent'}
                for n in names]

    def test_a_second_run_gets_its_own_budget(self):
        gr.MAX_REQUESTS = 2
        with mock.patch.object(gr, '_generate',
                               return_value=SimpleNamespace(text='{}', candidates=[])):
            app.attach_gemini_comparison(self._rows(['A', 'B']))
            first = gr.stats()
            app.attach_gemini_comparison(self._rows(['C', 'D']))
            second = gr.stats()
        self.assertEqual(first['requests'], 2)
        self.assertEqual(first['capped'], 0, 'the first run must fit in its budget')
        self.assertEqual(second['requests'], 2, 'counters must restart per run')
        self.assertEqual(second['capped'], 0,
                         'a fresh run must not inherit the previous run\'s spend')

    def test_the_cap_still_bites_within_one_run(self):
        gr.MAX_REQUESTS = 2
        with mock.patch.object(gr, '_generate',
                               return_value=SimpleNamespace(text='{}', candidates=[])):
            app.attach_gemini_comparison(self._rows(['A', 'B', 'C', 'D']))
        st = gr.stats()
        self.assertEqual(st['requests'], 2)
        self.assertEqual(st['capped'], 2, 'titles over the cap must be counted')

    def test_run_notes_report_this_run_not_the_process(self):
        gr.MAX_REQUESTS = 0  # unlimited
        with mock.patch.object(gr, '_generate',
                               return_value=SimpleNamespace(text='{}', candidates=[])):
            app.attach_gemini_comparison(self._rows(['A', 'B', 'C']))
            app.attach_gemini_comparison(self._rows(['D']))
        notes = {n['item']: n['value'] for n in app.gemini_run_notes([])}
        self.assertEqual(notes['requests made'], 1,
                         'the notes must describe the run they are printed beside')

    def test_cost_is_named_as_list_price_for_this_run(self):
        st = gr.stats()
        self.assertIn('run_search_cost_usd_at_list_price', st)
        self.assertEqual(st['est_search_cost_usd'],
                         st['run_search_cost_usd_at_list_price'])

    def test_the_entity_cache_survives_a_reset(self):
        """Resetting counters must not throw away paid-for answers."""
        with mock.patch.object(gr, '_generate',
                               return_value=SimpleNamespace(text='{}', candidates=[])):
            app.attach_gemini_comparison(self._rows(['Repeated Title']))
            app.attach_gemini_comparison(self._rows(['Repeated Title']))
            st = gr.stats()
        self.assertEqual(st['requests'], 0, 'the second run must not re-ask')
        self.assertEqual(st['cached'], 1, 'it must come from the cache instead')
