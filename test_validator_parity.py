"""Validator <-> Review parity on the deterministic ingest rules.

The Validator used to be a Movies/TV-shaped rule set: a Publisher or brand
file the Review covered in findings passed it clean. These cases lock in that
the two sections now flag the same columns on the same rows, for every schema
family, using the shared rules in ingest_rules.py.

Only the OFFLINE rules are compared -- the Review's discovery-based checks
(is this really the title's IMDb id / network?) are Review-only by design, so
every Review run here uses auto_fetch=False.
"""
import io
import sys

import openpyxl
import pandas as pd

import app
import validator

FAIL = []


def check(label, got, want):
    if got != want:
        FAIL.append("%s\n     got:  %r\n     want: %r" % (label, got, want))
    else:
        print("  ok  %s" % label)


class Up:
    def __init__(self, df, name='book.xlsx'):
        b = io.BytesIO()
        df.to_excel(b, index=False)
        self._b = b.getvalue()
        self.filename = name

    def read(self):
        return self._b


def validate(rows):
    """{column: severity} from the Validator."""
    _, s = validator.validate_workbook(Up(pd.DataFrame(rows)))
    return {f['column'].lower(): f['severity'] for f in s['failures']}


def review(rows):
    """{column: suggested} from the Review (offline)."""
    b = io.BytesIO()
    pd.DataFrame(rows).to_excel(b, index=False)
    xl, _ = app.build_review((b.getvalue(), 'book.xlsx'), auto_fetch=False)
    ws = openpyxl.load_workbook(io.BytesIO(xl))['Findings']
    hdr = [str(c.value).lower() for c in ws[1]]
    ci, si = hdr.index('column'), hdr.index('suggested value')
    return {str(r[ci]).lower(): str(r[si]) for r in ws.iter_rows(min_row=2, values_only=True)}


# ---------------------------------------------------------------- Publishers
print("\n-- Publishers (the Validator had NO rules for these) --")
BAD_PUB = dict(
    title='Vogue', title_category='Publishers',
    title_sub_category='Fashion & Beauty', companies='Conde Nast',
    brand_set='Competitive View', twitter_handle='@voguemagazine',
    instagram_user='https://instagram.com/vogue', facebook_page='vogue',
    linkedin_page='https://linkedin.com/company/vogue',
    genre='Fashion', imdb_id='http://www.imdb.com/title/tt1234567',
    twitter_search_terms='#vogue')
v, r = validate([BAD_PUB]), review([BAD_PUB])
for col in ('title', 'companies', 'brand_set', 'title_sub_category',
            'twitter_handle', 'instagram_user', 'genre', 'imdb_id'):
    check("both flag %s" % col, (col in v, col in r), (True, True))
check("Validator calls the publisher findings failures",
      v.get('companies'), validator.SEV_FAIL)

GOOD_PUB = dict(
    title='Vogue - DAR', title_category='Publishers',
    title_sub_category='Publication Type - Fashion & Beauty',
    companies='Pristine Brand',
    brand_set='LF // Publishing\nPristine DAR Brands',
    twitter_handle='voguemagazine', instagram_user='vogue',
    facebook_page='http://www.facebook.com/Vogue',
    twitter_search_terms='@voguemagazine|DAR|DAR\n#vogue|DAR|DAR')
check("a clean publisher row raises nothing", validate([GOOD_PUB]), {})

# ------------------------------------------------------------ brand schemas
print("\n-- Beauty / Beverages / Sports / General --")
check("the brand-schema rules are loaded", validator.TFX_OK, True)
check("all four schemas present",
      sorted(validator._TFX_RULES.get('schemas', {})),
      ['beauty', 'beverages', 'general', 'sports'])
BAD_BEAUTY = dict(title='Glossier', title_category='Health & Beauty',
                  perspective='Standard', title_sub_category='', companies='',
                  brand_set='', active='t', twitter_search_terms='#glossier')
v, r = validate([BAD_BEAUTY]), review([BAD_BEAUTY])
check("both flag companies on a Standard beauty row",
      ('companies' in v, 'companies' in r), (True, True))
check("both flag the missing ' - DAR' title",
      ('title' in v, 'title' in r), (True, True))

# ------------------------------------------------------ Movies/TV parity gaps
print("\n-- Movies/TV rules the Validator was missing --")
MOV = dict(title='Some Film', title_category='Movies', released_on='2026-01-05',
           genre='Action\nDrama', primary_genre='Comedy', companies='Neon',
           title_sub_category='Release - Wide\nStudio - Independent')
check("primary_genre outside the row's own genres fails",
      validate([MOV]).get('primary_genre'), validator.SEV_FAIL)
check("...and the Review flags it too", 'primary_genre' in review([MOV]), True)

MOV2 = dict(MOV, primary_genre='Action',
            title_sub_category='Release - Wide\nStudio - Independent\nLanguage Type - English')
check("a Language Type line on a Movie fails",
      validate([MOV2]).get('title_sub_category'), validator.SEV_FAIL)

MOV3 = dict(MOV, primary_genre='Action', twitter_handle='http://twitter.com/somefilm')
check("a URL in twitter_handle fails",
      validate([MOV3]).get('twitter_handle'), validator.SEV_FAIL)

MOV4 = dict(MOV, primary_genre='Action', twitter_handle='somefilm|2026-03-13')
check("an attribution window is NOT mistaken for a bad handle",
      validate([MOV4]).get('twitter_handle'), None)

MOV5 = dict(MOV, primary_genre='Action', instagram_user='somefilmfanpage')
check("a Fanpage instagram handle fails",
      validate([MOV5]).get('instagram_user'), validator.SEV_FAIL)

MOV6 = dict(MOV, primary_genre='Action', companies='')
check("a blank companies on a non-DAR row warns",
      validate([MOV6]).get('companies'), validator.SEV_WARN)

MOV7 = dict(MOV, primary_genre='Action', facebook_verified='yes')
check("a malformed verified cell fails",
      validate([MOV7]).get('facebook_verified'), validator.SEV_FAIL)
MOV8 = dict(MOV, primary_genre='Action',
            facebook_verified='TRUE|http://www.facebook.com/somefilm')
check("a well-formed verified cell passes",
      validate([MOV8]).get('facebook_verified'), None)

print("\n-- Talent / Video Games --")
TAL = dict(title='Jane Doe - DAR', title_category='Talent', companies='Someone Else',
           title_sub_category='Actor', brand_set='LF // Talent')
check("a DAR talent row must be Pristine Brand",
      validate([TAL]).get('companies'), validator.SEV_FAIL)
check("an unlabelled talent sub-category line fails",
      validate([TAL]).get('title_sub_category'), validator.SEV_FAIL)
GAME = dict(title='Akatori - DAR', title_category='Video Game',
            companies='Pristine Brand', brand_set='LF // Video Games // Games',
            title_sub_category='Platform - PC\nDeveloper - Studio X')
check("a well-formed game sub-category passes",
      validate([GAME]).get('title_sub_category'), None)

print("\n-- regression: the existing rules still fire --")
check("blank title still fails",
      validate([dict(title='', title_category='Movies')]).get('title'),
      validator.SEV_FAIL)
check("a bad Rotten Tomatoes URL still fails",
      validate([dict(MOV, primary_genre='Action',
                     rottentomatoes='http://www.rottentomatoes.com/tv/some_film')
                ]).get('rottentomatoes'), validator.SEV_FAIL)

print()
if FAIL:
    print("FAILED (%d):" % len(FAIL))
    for x in FAIL:
        print("  x  " + x)
    sys.exit(1)
print("Validator/Review parity checks passed.")
